"""Import the Plett tourism xlsx into data/activities.json."""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "Copy of Plettenberg Bay Comprehensive Tourism Guide V2 (1).xlsx"
OUT = ROOT / "data" / "activities.json"


def slugify(*parts: str) -> str:
    text = "-".join(p for p in parts if p)
    text = re.sub(r"[^a-z0-9]+", "-", text.lower())
    return text.strip("-")[:80] or "activity"


def category_for(company: str, name: str) -> str:
    hay = f"{company} {name}".lower()
    if any(word in hay for word in ("kloof", "canyon", "quad", "water park", "adventure land")):
        return "Adventure"
    if "golf" in hay and "padel" not in hay:
        return "Golf"
    if "padel" in hay:
        return "Padel"
    if any(word in hay for word in ("wine", "vineyard", "mcc", "cap classique")):
        return "Wine"
    if any(word in hay for word in ("hike", "hiking", "trail", "bike", "sanparks", "capenature", "cairnbrogie")):
        return "Trails"
    return "Experiences"


def cell(row, index):
    if index >= len(row) or row[index] is None:
        return ""
    return row[index]


def main() -> None:
    wb = load_workbook(XLSX, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    activities = []
    seen = set()

    for row in rows[1:]:
        company = str(cell(row, 0)).strip()
        name = str(cell(row, 1)).strip()
        if not company or not name:
            continue

        price_raw = cell(row, 2)
        if isinstance(price_raw, (int, float)):
            price_numeric = float(price_raw)
            price_label = f"R {price_numeric:,.0f}" if price_numeric == int(price_numeric) else f"R {price_numeric:,.2f}"
        else:
            price_numeric = None
            price_label = str(price_raw).strip() or "Variable"

        lat_raw = str(cell(row, 6)).replace("\\", "").strip()
        try:
            lat = float(lat_raw)
        except ValueError:
            lat = None

        lng_raw = cell(row, 7)
        image_parts = []
        if isinstance(lng_raw, (int, float)):
            lng = float(lng_raw)
            image_parts = [str(cell(row, i)).strip() for i in range(8, 11) if str(cell(row, i)).strip()]
        else:
            lng = None
            image_parts = [str(x).strip() for x in (lng_raw, cell(row, 8), cell(row, 9), cell(row, 10)) if str(x).strip()]

        image_url = ",".join(image_parts)

        slug = slugify(company, name)
        original = slug
        n = 2
        while slug in seen:
            slug = f"{original}-{n}"
            n += 1
        seen.add(slug)

        activities.append({
            "id": slug,
            "company": company,
            "name": name,
            "category": category_for(company, name),
            "price_label": price_label,
            "price_numeric": price_numeric,
            "phone": str(cell(row, 3)).strip(),
            "email": str(cell(row, 4)).strip(),
            "booking_email": str(cell(row, 4)).strip(),
            "booking_url": str(cell(row, 5)).strip(),
            "lat": lat,
            "lng": lng,
            "image_url": image_url,
            "location": "Plettenberg Bay",
            "description": "",
            "published": True,
        })

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(activities, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {len(activities)} activities to {OUT}")


if __name__ == "__main__":
    main()
